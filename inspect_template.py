from pptx import Presentation

prs = Presentation('Report/Final Presentation-Capstone Project in Data Science II.pptx')
print(f'Slides: {len(prs.slides)}')
print(f'Slide size: {prs.slide_width.inches:.2f}" x {prs.slide_height.inches:.2f}"')
print()

for i, sl in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} (layout={sl.slide_layout.name}) ===')
    for j, sh in enumerate(sl.shapes):
        l = round(sh.left/914400, 2)
        t = round(sh.top/914400, 2)
        w = round(sh.width/914400, 2)
        h = round(sh.height/914400, 2)
        print(f'  [{j}] type={sh.shape_type} name={repr(sh.name)} pos=({l},{t}) size={w}x{h}')
        if sh.has_text_frame:
            for k, p in enumerate(sh.text_frame.paragraphs):
                txt = p.text.strip()
                if txt:
                    print(f'       para[{k}]: {repr(txt[:90])}')
        if sh.shape_type == 13:
            print(f'       (picture/image)')
    print()
